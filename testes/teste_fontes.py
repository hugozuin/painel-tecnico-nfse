"""Confere o JSON gerado contra as planilhas, lendo-as por outro caminho de
código (mapa de mesclagens em vez de grade preenchida)."""
import json
import re
import sys
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent

BASE = RAIZ / "fontes" / "nacional"
falhas = 0

def conferir(titulo, condicao, extra=""):
    global falhas
    print(f"  {'ok  ' if condicao else 'FALHA'} {titulo} {'' if condicao else extra}")
    if not condicao:
        falhas += 1

def mapa_de_mesclagens(aba):
    ancora = {}
    for faixa in aba.merged_cells.ranges:
        for linha in range(faixa.min_row, faixa.max_row + 1):
            for coluna in range(faixa.min_col, faixa.max_col + 1):
                ancora[(linha, coluna)] = (faixa.min_row, faixa.min_col)
    return lambda linha, coluna: aba.cell(*ancora.get((linha, coluna), (linha, coluna))).value

def limpo(valor):
    return re.sub(r"\s+", " ", str(valor or "").replace("\xa0", " ")).strip()

ibscbs = json.load(open(RAIZ / "definicoes" / "ibscbs.json", encoding="utf-8"))
de_para = json.load(open(RAIZ / "definicoes" / "de-para-nacional.json", encoding="utf-8"))

print("\n== anexo VIII, correlação ==")
livro = openpyxl.load_workbook(BASE / "anexoviii-correlacaoitemnbsindopcclasstrib_ibscbs_v1-01-00.xlsx")
aba = livro["tabela geral"]
ler = mapa_de_mesclagens(aba)
esperadas = set()
for linha in range(2, aba.max_row + 1):
    item, nbs = limpo(ler(linha, 1)), limpo(ler(linha, 3))
    if not item and not nbs:
        continue
    indop = limpo(ler(linha, 7))
    indop = indop.split(".")[0].zfill(6) if indop else ""
    classificacao = limpo(ler(linha, 9))
    classificacao = classificacao.zfill(6) if classificacao.isdigit() else classificacao
    esperadas.add((item, nbs, limpo(ler(linha, 5)), limpo(ler(linha, 6)), indop, limpo(ler(linha, 8)), classificacao))
geradas = {(r[0], r[1], r[2], r[3], r[4], limpo(ibscbs["locais"][r[5]]) if r[5] >= 0 else "", r[6]) for r in ibscbs["relacoes"]}
conferir("mesmo conjunto de relações", esperadas == geradas,
         f"faltam {len(esperadas - geradas)}, sobram {len(geradas - esperadas)}")
conferir("quantidade de relações", len(esperadas) == len(ibscbs["relacoes"]), f"{len(esperadas)} x {len(ibscbs['relacoes'])}")
amostra = {(r[0], r[1], r[6]) for r in esperadas if r[0] == "01.01" and r[1] == "1.1502.90.00"}
conferir("NBS mesclado recebe as três classificações", {c for _, _, c in amostra} == {"000001", "200043", "200044"}, str(amostra))

for linha in range(2, aba.max_row + 1):
    codigo = limpo(ler(linha, 9))
    if codigo:
        codigo = codigo.zfill(6) if codigo.isdigit() else codigo
        if limpo(ibscbs["cClassTrib"].get(codigo)) != limpo(ler(linha, 10)):
            conferir(f"nome do cClassTrib {codigo}", False)
            break
else:
    conferir("nomes de todos os cClassTrib iguais à planilha", True)

print("\n== anexo VII, indOp ==")
livro = openpyxl.load_workbook(BASE / "anexovii-indop_ibscbs_v1-02-00.xlsx")
aba = livro["cIndOp Public"]
codigos = {}
for linha in aba.iter_rows(min_row=2, values_only=True):
    if linha[0]:
        codigos[str(linha[0]).strip().zfill(6)] = limpo(linha[1])
conferir("mesmos 39 códigos", set(codigos) == set(ibscbs["indOp"]) and len(codigos) == 39, str(len(codigos)))
conferir("mesmo tipo de operação em todos", all(limpo(ibscbs["indOp"][c]["tipoOperacao"]) == t for c, t in codigos.items()))

print("\n== anexo VI, leiaute e regras ==")
livro = openpyxl.load_workbook(BASE / "anexovi-leiautesrn_rtc_ibscbs-v1-04-00-2013-nt009.xlsx")
aba = livro["LEIAUTE DPS_NFS-e - RT"]
ler = mapa_de_mesclagens(aba)
por_chave = {e["caminho"] + e["tag"]: e for e in de_para["entradas"]}
linhas_com_campo = 0
divergentes = []
for linha in range(2, aba.max_row + 1):
    campo = limpo(ler(linha, 3))
    if not campo or campo == "-":
        continue
    linhas_com_campo += 1
    chave = limpo(ler(linha, 2)).replace(" ", "")
    chave = (chave if chave.endswith("/") else chave + "/") + campo
    entrada = por_chave.get(chave)
    descricao = limpo(ler(linha, 8))
    if entrada is None or limpo(entrada["descricao"]) != ("" if descricao == "-" else descricao):
        divergentes.append(chave)
conferir("todas as tags do leiaute presentes com a descrição da planilha", not divergentes, str(divergentes[:5]))
conferir("mesma quantidade de tags", linhas_com_campo == len(de_para["entradas"]), f"{linhas_com_campo} x {len(de_para['entradas'])}")

aba = livro["RN DPS_NFS-e"]
ler = mapa_de_mesclagens(aba)
regras_planilha = set()
caminho = campo = ""
for linha in range(4, aba.max_row + 1):
    if limpo(ler(linha, 2)) not in ("", "-") or limpo(ler(linha, 3)) not in ("", "-"):
        caminho, campo = limpo(ler(linha, 2)).replace(" ", ""), limpo(ler(linha, 3))
    texto = limpo(ler(linha, 4))
    if texto in ("", "-") or not campo:
        continue
    codigo = limpo(ler(linha, 8))
    regras_planilha.add((codigo if codigo != "-" else "", texto))
regras_json = {(r["codigo"], limpo(r["regra"])) for e in de_para["entradas"] for r in e["regras"]}
conferir("todas as regras da planilha estão no de-para", regras_planilha <= regras_json,
         f"faltam {len(regras_planilha - regras_json)}")
conferir("nenhuma regra inventada", regras_json <= regras_planilha, f"sobram {len(regras_json - regras_planilha)}")

texto_json = (RAIZ / "definicoes" / "de-para-nacional.json").read_text(encoding="utf-8")
conferir("sem trecho de código no de-para publicado",
         not any(marca in texto_json for marca in ("CampoTecno(", "SetarCampo", "=>", "formatRounding(")))

print("\nConferência contra as planilhas passou." if falhas == 0 else f"\n{falhas} conferência(s) falharam.")
sys.exit(1 if falhas else 0)
