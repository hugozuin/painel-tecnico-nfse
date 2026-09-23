"""Gera os arquivos de definição do de-para e da relação IBS/CBS.

Todo o conteúdo sai de documentos, sem texto escrito à mão:

  Anexo VI    leiaute da DPS/NFS-e e regras de negócio
  Anexo VII   tabela de indicadores de operação (indOp)
  Anexo VIII  correlação entre item da LC 116, NBS, indOp e cClassTrib
  Mapping.txt campo do dataset do componente para o caminho no XML
  LoadEnvio   campo do TX2 para o campo do dataset (script do padrão Nacional)
  lib props   campo do JSON do PlugNotas para o campo do TX2

Os anexos são públicos e ficam em fontes/nacional. Script, mapeamento e lib
são internos: entram só por parâmetro e não vão para o repositório. O JSON
gerado guarda nomes de campo e números de linha, nunca trechos de código.

Uso:
  python ferramentas/gerar_definicoes.py --anexos fontes/nacional \\
      --script <LoadEnvio.txt> --mapeamento <Mapping.txt> --lib <pasta props>

Sem script, mapeamento e lib, o de-para sai só com o lado do Nacional.
Requer Python 3.10+ e openpyxl.
"""

import argparse
import difflib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ANEXO_VI = "anexovi-leiautesrn_rtc_ibscbs-v1-04-00-2013-nt009.xlsx"
ANEXO_VII = "anexovii-indop_ibscbs_v1-02-00.xlsx"
ANEXO_VIII = "anexoviii-correlacaoitemnbsindopcclasstrib_ibscbs_v1-01-00.xlsx"
PREFIXO_NFSE = "NFSe/infNFSe/"

# Raízes do JSON confirmadas pelo exemplo de emissão da coleção do Postman.
RAIZES_JSON_CONFIRMADAS = {
    "prestador": "prestador",
    "tomador": "tomador",
    "servico": "servico[]",
    "ibscbs": "ibscbs",
}

# Gravações que copiam o valor sem conversão quando recebem CampoTecno puro.
# As de moeda convertem o formato numérico, então não contam como cópia direta.
SETTERS = {
    "SetarCampoValorTamanhoObrigatorio": (0, 1, True),
    "SetarCampoValorTamanho": (0, 1, True),
    "SetarCampoValorCurrency": (0, 1, False),
    "SetarCampoValor": (0, 1, True),
    "SetarCampoTamanhoMinimo": (1, None, False),
}


def texto(valor):
    if valor is None:
        return ""
    conteudo = str(valor).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    return "\n".join(linha.rstrip() for linha in conteudo.split("\n")).strip()


def vazio(valor):
    return texto(valor) in ("", "-")


def grade_com_mesclagens(aba):
    """Lê a aba replicando o valor de cada mesclagem em todas as células dela."""
    grade = [[celula.value for celula in linha] for linha in aba.iter_rows()]
    for faixa in aba.merged_cells.ranges:
        origem = grade[faixa.min_row - 1][faixa.min_col - 1]
        for linha in range(faixa.min_row - 1, faixa.max_row):
            for coluna in range(faixa.min_col - 1, faixa.max_col):
                if linha < len(grade) and coluna < len(grade[linha]):
                    grade[linha][coluna] = origem
    return grade


def caminho_normalizado(caminho):
    limpo = texto(caminho).replace(" ", "")
    if limpo and not limpo.endswith("/"):
        limpo += "/"
    return limpo


def titulo_da_descricao(descricao):
    """Título curto tirado da própria descrição: o trecho antes dos dois
    pontos ou a primeira frase. Não reescreve nada do anexo."""
    corrido = re.sub(r"\s+", " ", descricao).strip()
    if ":" in corrido:
        corrido = corrido.split(":", 1)[0]
    else:
        sentenca = re.match(r"(.+?\.)(\s|$)", corrido)
        if sentenca:
            corrido = sentenca.group(1)
    corrido = corrido.strip().rstrip(" .;")
    return corrido if len(corrido) <= 120 else corrido[:117].rstrip() + "…"


# ------------------------------------------------------------------ Anexo VI

def ler_leiaute(pasta_anexos):
    livro = openpyxl.load_workbook(pasta_anexos / ANEXO_VI, data_only=True)
    grade = grade_com_mesclagens(livro["LEIAUTE DPS_NFS-e - RT"])
    entradas = []
    for linha in grade[1:]:
        campo = texto(linha[2])
        if not campo or campo == "-":
            continue
        descricao = texto(linha[7])
        entradas.append({
            "id": texto(linha[0]),
            "tag": campo,
            "caminho": caminho_normalizado(linha[1]),
            "elemento": texto(linha[3]),
            "tipo": texto(linha[4]),
            "ocorrencia": texto(linha[5]),
            "tamanho": texto(linha[6]),
            "titulo": campo if vazio(descricao) else titulo_da_descricao(descricao),
            "descricao": "" if vazio(descricao) else descricao,
            "notas": "" if vazio(linha[8]) else texto(linha[8]),
            "regras": [],
        })
    return entradas


def ler_regras(pasta_anexos):
    livro = openpyxl.load_workbook(pasta_anexos / ANEXO_VI, data_only=True)
    grade = grade_com_mesclagens(livro["RN DPS_NFS-e"])
    regras = {}
    caminho_atual = campo_atual = ""
    for linha in grade[3:]:
        if not vazio(linha[1]) or not vazio(linha[2]):
            caminho_atual = caminho_normalizado(linha[1])
            campo_atual = texto(linha[2])
        regra = texto(linha[3])
        if vazio(regra) or not campo_atual:
            continue
        registro = {
            "codigo": "" if vazio(linha[7]) else texto(linha[7]),
            "regra": regra,
            "mensagem": "" if vazio(linha[8]) else texto(linha[8]),
            "aplicacao": "" if vazio(linha[5]) else texto(linha[5]),
            "efeito": "" if vazio(linha[6]) else texto(linha[6]),
            "nivel": "" if vazio(linha[9]) else texto(linha[9]),
            "observacao": "" if len(linha) < 15 or vazio(linha[14]) else texto(linha[14]),
        }
        lista = regras.setdefault(caminho_atual + campo_atual, [])
        if registro not in lista:
            lista.append(registro)
    return regras


def aproximar_caminho(chave, entradas):
    """Casa um caminho da aba de regras com o leiaute quando os dois diferem
    na grafia, como totalTrib e totTrib. Exige a mesma tag e uma semelhança
    alta e sem empate."""
    tag = chave.rstrip("/").rsplit("/", 1)[-1]
    candidatos = sorted(
        ((difflib.SequenceMatcher(None, chave, e["caminho"] + e["tag"]).ratio(), indice)
         for indice, e in enumerate(entradas) if e["tag"] == tag),
        reverse=True,
    )
    if not candidatos or candidatos[0][0] < 0.75:
        return None
    if len(candidatos) > 1 and candidatos[0][0] - candidatos[1][0] < 0.05:
        return None
    return entradas[candidatos[0][1]]


# ------------------------------------------------------------------ Mapping

def ler_mapeamento(arquivo):
    campos = {}
    secao = ""
    for linha in arquivo.read_text(encoding="latin-1").splitlines():
        limpa = linha.strip()
        if limpa.startswith("["):
            secao = limpa.strip("[]")
            continue
        if secao != "DPS" or not limpa or limpa.startswith(";") or "=" not in limpa:
            continue
        chave, caminho = limpa.split("=", 1)
        partes = caminho.strip().replace("@", "").split("/")
        campos[chave.strip()] = PREFIXO_NFSE + "/".join(partes[:-1]) + "/" + partes[-1]
    return campos


# ------------------------------------------------------------------ script Pascal

def sem_comentarios(linha):
    resultado = []
    em_texto = False
    posicao = 0
    while posicao < len(linha):
        caractere = linha[posicao]
        if caractere == "'":
            em_texto = not em_texto
        elif not em_texto and linha.startswith("//", posicao):
            break
        resultado.append(caractere)
        posicao += 1
    return "".join(resultado)


def argumentos_da_chamada(conteudo, inicio):
    nivel = 0
    em_texto = False
    atual = []
    argumentos = []
    for posicao in range(inicio, len(conteudo)):
        caractere = conteudo[posicao]
        if caractere == "'":
            em_texto = not em_texto
        if not em_texto:
            if caractere == "(":
                nivel += 1
                if nivel == 1:
                    continue
            elif caractere == ")":
                nivel -= 1
                if nivel == 0:
                    argumentos.append("".join(atual).strip())
                    return argumentos
            elif caractere == "," and nivel == 1:
                argumentos.append("".join(atual).strip())
                atual = []
                continue
        atual.append(caractere)
    return argumentos


def campos_tx2(expressao):
    return re.findall(r"CampoTecno\w*\('(\w+)'\)", expressao)


def ler_script(arquivo):
    """Liga cada campo do dataset aos campos do TX2 que determinam seu valor.

      direto     o campo TX2 aparece na própria chamada que grava o dataset
      calculado  o valor gravado é uma variável montada a partir de campos TX2
      condicao   o valor gravado é fixo e depende do if ou case que o controla

    A análise considera os blocos begin, case e try para não ligar uma
    atribuição feita em outro ramo do mesmo case.
    """
    linhas = [sem_comentarios(linha) for linha in arquivo.read_text(encoding="latin-1").splitlines()]
    conteudo = "\n".join(linhas)
    sem_textos = re.sub(r"'[^']*'", lambda achado: " " * len(achado.group(0)), conteudo)
    inicio_linha = [0]
    for linha in linhas:
        inicio_linha.append(inicio_linha[-1] + len(linha) + 1)

    def numero_da_linha(posicao):
        baixo, alto = 0, len(inicio_linha) - 1
        while baixo < alto:
            meio = (baixo + alto + 1) // 2
            if inicio_linha[meio] <= posicao:
                baixo = meio
            else:
                alto = meio - 1
        return baixo + 1

    marcas = []
    pilha = []
    fim_do_bloco = {}
    for achado in re.finditer(r"\b(begin|case|try|end)\b", sem_textos, flags=re.IGNORECASE):
        palavra = achado.group(1).lower()
        if palavra == "end":
            if pilha:
                fim_do_bloco[pilha.pop()[1]] = achado.end()
        else:
            pilha.append((palavra, achado.start()))
        marcas.append((achado.start(), tuple(pilha)))

    def pilha_em(posicao):
        baixo, alto, resposta = 0, len(marcas) - 1, ()
        while baixo <= alto:
            meio = (baixo + alto) // 2
            if marcas[meio][0] < posicao:
                resposta = marcas[meio][1]
                baixo = meio + 1
            else:
                alto = meio - 1
        return resposta

    def alcanca(posicao_atribuicao, posicao_uso):
        origem = pilha_em(posicao_atribuicao)
        destino = pilha_em(posicao_uso)
        comum = 0
        while comum < min(len(origem), len(destino)) and origem[comum] == destino[comum]:
            comum += 1
        if comum == len(origem) or comum >= len(destino):
            return True
        if comum > 0 and origem[comum - 1][0] == "case":
            return False
        bloco_origem, bloco_destino = origem[comum], destino[comum]
        if bloco_origem[0] == "begin" and bloco_destino[0] == "begin":
            fim = fim_do_bloco.get(bloco_origem[1])
            if fim is not None and re.fullmatch(r"\s*;?\s*else\s*", sem_textos[fim:bloco_destino[1]]):
                return False
        return True

    def envolve(posicao_atribuicao, posicao_uso):
        origem = pilha_em(posicao_atribuicao)
        return pilha_em(posicao_uso)[:len(origem)] == origem

    def guarda(posicao):
        """Condição de uma atribuição de linha única (if/then, else ou rótulo
        de case). None quando a atribuição não é condicional."""
        numero = numero_da_linha(posicao)
        antes = sem_textos[inicio_linha[numero - 1]:posicao]
        anterior = ""
        for indice in range(numero - 2, max(numero - 6, -1), -1):
            if linhas[indice].strip():
                anterior = linhas[indice]
                break
        mesma_linha = re.search(r"\bif\b(.*)\bthen\b", antes)
        if mesma_linha:
            return mesma_linha.group(1)
        if re.search(r"\belse\b", antes) or re.match(r"\s*[\d', ]+:(?!=)", antes):
            return ""
        linha_anterior = re.search(r"\bif\b(.*)\bthen\s*$", anterior)
        if linha_anterior:
            return linha_anterior.group(1)
        if re.search(r"\belse\s*$", anterior) or re.match(r"\s*[\d', ]+:\s*$", anterior):
            return ""
        return None

    # Uma atribuição termina no ponto e vírgula ou antes de else/end, como em
    # "if X then v := A else v := B;".
    atribuicoes = [
        (achado.start(), achado.group(1), achado.end())
        for achado in re.finditer(r"\b(\w+)\s*:=\s*(?:(?!\belse\b|\bend\b)[^;])*", sem_textos, flags=re.IGNORECASE)
    ]

    def campos_do_trecho(trecho, posicao):
        campos = campos_tx2(trecho)
        for variavel in re.findall(r"CampoTecno\w*\(\s*(_\w+)\s*\)", trecho):
            anteriores = [item for item in atribuicoes if item[1] == variavel and item[0] < posicao]
            if anteriores:
                literal = re.search(r":=\s*'(\w+)'", conteudo[anteriores[-1][0]:anteriores[-1][2]])
                if literal:
                    campos.append(literal.group(1))
        return campos

    def origem_da_variavel(variavel, antes_de, profundidade=0):
        alcancam = [
            item for item in atribuicoes
            if item[1] == variavel and item[0] < antes_de and alcanca(item[0], antes_de)
        ]
        dominante = None
        for indice in range(len(alcancam) - 1, -1, -1):
            atual = alcancam[indice]
            if not envolve(atual[0], antes_de):
                continue
            if guarda(atual[0]) is None:
                dominante = indice
                break
            if indice > 0:
                anterior = alcancam[indice - 1]
                entre = sem_textos[anterior[2]:atual[0]]
                if (re.fullmatch(r"\s*;?\s*else\s*", entre) and envolve(anterior[0], antes_de)
                        and guarda(anterior[0])):
                    dominante = indice - 1
                    break
        escolhidas = alcancam[dominante:] if dominante is not None else alcancam

        campos = []
        for inicio, _, fim in escolhidas:
            trecho = conteudo[inicio:fim]
            campos.extend(campos_do_trecho(trecho, inicio))
            condicao = guarda(inicio)
            if condicao:
                campos.extend(campos_tx2(condicao))
            if condicao is not None or profundidade >= 3:
                continue
            citadas = set(re.findall(r"\b(_\w+|v[A-Z]\w*)\b", trecho.split(":=", 1)[-1]))
            for outra in sorted(citadas - {variavel}):
                if re.search(rf"CampoTecno\w*\(\s*{re.escape(outra)}\s*\)", trecho):
                    continue
                campos.extend(origem_da_variavel(outra, inicio, profundidade + 1))
        return list(dict.fromkeys(campos))

    def cabecalho_envolvente(posicao):
        numero = numero_da_linha(posicao)
        recuo = len(linhas[numero - 1]) - len(linhas[numero - 1].lstrip())
        for anterior in range(numero - 1, max(numero - 40, 0), -1):
            linha = linhas[anterior - 1]
            if not linha.strip():
                continue
            recuo_anterior = len(linha) - len(linha.lstrip())
            if recuo_anterior >= recuo:
                continue
            achado = re.search(r"\b(if|case)\b(.*?)\b(then|of)\b", linha)
            if achado:
                return achado.group(2), inicio_linha[anterior - 1]
            if not re.match(r"\s*(begin|end|else|[\d, ]+:)", linha):
                return "", 0
            recuo = recuo_anterior
        return "", 0

    padrao_setter = re.compile(r"\b(" + "|".join(sorted(SETTERS, key=len, reverse=True)) + r")\s*\(")
    resultado = {}
    for achado in padrao_setter.finditer(conteudo):
        funcao = achado.group(1)
        argumentos = argumentos_da_chamada(conteudo, achado.end() - 1)
        posicao_campo, posicao_valor, copia_simples = SETTERS[funcao]
        if len(argumentos) <= posicao_campo:
            continue
        nome = re.match(r"'(\w+)'", argumentos[posicao_campo])
        if not nome:
            continue
        fontes = []
        if funcao == "SetarCampoTamanhoMinimo":
            tx2 = re.match(r"'(\w+)'", argumentos[0])
            if tx2:
                fontes.append(("direto", tx2.group(1), False))
        else:
            valor = argumentos[posicao_valor] if len(argumentos) > posicao_valor else ""
            puro = bool(re.fullmatch(r"CampoTecno\('\w+'\)", valor)) and copia_simples
            fontes.extend(("direto", campo, puro) for campo in campos_tx2(valor))
            for variavel in re.findall(r"\b(_\w+|v[A-Z]\w*)\b", valor):
                fontes.extend(("calculado", campo, False) for campo in origem_da_variavel(variavel, achado.start()))
            if not fontes:
                condicao, posicao_condicao = cabecalho_envolvente(achado.start())
                fontes.extend(("condicao", campo, False) for campo in campos_tx2(condicao))
                for variavel in re.findall(r"\b(_\w+|v[A-Z]\w*)\b", condicao):
                    fontes.extend(("condicao", campo, False) for campo in origem_da_variavel(variavel, posicao_condicao))

        registro = resultado.setdefault(nome.group(1), {"linhas": [], "fontes": {}})
        registro["linhas"].append(numero_da_linha(achado.start()))
        for uso, campo, puro in fontes:
            anterior = registro["fontes"].get(campo)
            if anterior is None:
                registro["fontes"][campo] = {"uso": uso, "copiaDireta": puro}
            else:
                anterior["copiaDireta"] = anterior["copiaDireta"] and puro
    return resultado


# ------------------------------------------------------------------ lib do PlugNotas

def ler_parametros_do_servico(pasta_lib):
    arquivo = pasta_lib / "servico" / "getServicoProps.js"
    parametros = {}
    if not arquivo.exists():
        return parametros
    for funcao, argumento in re.findall(r"\b(get\w+Props)\(\s*([\w.?]+)", arquivo.read_text(encoding="utf-8")):
        partes = argumento.replace("?", "").split(".")
        if partes[0] == "servico":
            parametros[funcao] = "servico[]" + ("." + ".".join(partes[1:]) if len(partes) > 1 else "")
    return parametros


def pares_maiusculos(conteudo):
    pares = []
    for achado in re.finditer(r"(?:^|[{,\s])([A-Z][A-Za-z0-9_]*)\s*:", conteudo):
        posicao = achado.end()
        nivel = 0
        em_texto = None
        expressao = []
        while posicao < len(conteudo):
            caractere = conteudo[posicao]
            if em_texto:
                if caractere == em_texto and conteudo[posicao - 1] != "\\":
                    em_texto = None
            elif caractere in "'\"`":
                em_texto = caractere
            elif caractere in "([{":
                nivel += 1
            elif caractere in ")]}":
                if nivel == 0:
                    break
                nivel -= 1
            elif caractere == "," and nivel == 0:
                break
            expressao.append(caractere)
            posicao += 1
        limpa = re.sub(r"\s+", " ", re.sub(r"//[^\n]*", "", "".join(expressao))).strip()
        if limpa:
            pares.append((achado.group(1), limpa))
    return pares


def caminhos_da_expressao(expressao, variaveis):
    encontrados = []
    for achado in re.finditer(r"\b([a-zA-Z_]\w*)((?:\??\.[A-Za-z_]\w*)+)", expressao):
        base = variaveis.get(achado.group(1))
        if not base:
            continue
        sufixo = achado.group(2).replace("?.", ".")
        encontrados.extend(caminho + sufixo for caminho in base)
    solto = re.fullmatch(r"\s*([a-zA-Z_]\w*)\s*", expressao)
    if solto and solto.group(1) in variaveis:
        encontrados.extend(variaveis[solto.group(1)])
    return list(dict.fromkeys(encontrados))


def ler_lib(pasta_lib):
    parametros_servico = ler_parametros_do_servico(pasta_lib)
    resultado = {}
    raizes_sem_confirmacao = set()

    for arquivo in sorted(pasta_lib.rglob("*.js")):
        if arquivo.name == "index.js":
            continue
        conteudo = arquivo.read_text(encoding="utf-8")
        conteudo = re.sub(r"/\*.*?\*/", "", conteudo, flags=re.DOTALL)
        conteudo = re.sub(r"(^|[^:])//[^\n]*", r"\1", conteudo)
        funcao = arquivo.stem
        assinatura = re.search(rf"const\s+{funcao}\s*=\s*\(?\s*([a-zA-Z_]\w*)", conteudo)
        variaveis = {}
        if assinatura:
            parametro = assinatura.group(1)
            if funcao in parametros_servico:
                variaveis[parametro] = [parametros_servico[funcao]]
            elif parametro in RAIZES_JSON_CONFIRMADAS:
                variaveis[parametro] = [RAIZES_JSON_CONFIRMADAS[parametro]]
            else:
                raizes_sem_confirmacao.add(f"{funcao}({parametro})")

        for _ in range(3):
            for nomes, expressao in re.findall(r"const\s*\{([^}]+)\}\s*=\s*([^\n]+)", conteudo):
                bases = caminhos_da_expressao(expressao.split("||")[0], variaveis)
                for nome in [parte.strip() for parte in nomes.split(",") if parte.strip()]:
                    if bases:
                        variaveis[nome] = [base + "." + nome for base in bases]
            for nome, expressao in re.findall(r"const\s+([a-zA-Z_]\w*)\s*=\s*([^\n]+)", conteudo):
                if nome in variaveis or "=>" in expressao:
                    continue
                if "?" in expressao and ":" in expressao:
                    bases = []
                    for alternativa in re.split(r"\?|:", expressao)[1:]:
                        bases.extend(caminhos_da_expressao(alternativa.strip(), variaveis))
                else:
                    bases = caminhos_da_expressao(expressao, variaveis)
                if bases:
                    variaveis[nome] = bases
            for base, iterador in re.findall(r"([a-zA-Z_][\w?.]*)\s*\??\.\s*map\(\s*\(?\s*(\w+)", conteudo):
                partes = base.replace("?", "").rstrip(".").split(".")
                origem = variaveis.get(partes[0])
                if origem:
                    variaveis[iterador] = [
                        caminho + ("." + ".".join(partes[1:]) if len(partes) > 1 else "") + "[]"
                        for caminho in origem
                    ]

        for campo, expressao in pares_maiusculos(conteudo):
            caminhos = caminhos_da_expressao(expressao, variaveis)
            copia = bool(re.fullmatch(r"[a-zA-Z_]\w*(\??\.\w+)+(\s*\|\|\s*'')?", expressao))
            registro = resultado.setdefault(campo, {"json": [], "copiaDireta": []})
            for caminho in caminhos:
                if caminho not in registro["json"]:
                    registro["json"].append(caminho)
                if copia and caminho not in registro["copiaDireta"]:
                    registro["copiaDireta"].append(caminho)
    return resultado, sorted(raizes_sem_confirmacao)


# ------------------------------------------------------------------ de-para

def gerar_de_para(pasta_anexos, script, mapeamento, lib):
    entradas = ler_leiaute(pasta_anexos)
    regras = ler_regras(pasta_anexos)
    por_chave = {e["caminho"] + e["tag"]: e for e in entradas}
    aproximadas = []
    sem_tag = []

    for chave, lista in regras.items():
        alvo = por_chave.get(chave)
        aproximado = alvo is None
        if alvo is None:
            alvo = aproximar_caminho(chave, entradas)
        if alvo is None:
            sem_tag.append(chave)
            continue
        if aproximado:
            aproximadas.append(f"{chave} -> {alvo['caminho']}{alvo['tag']}")
        for regra in lista:
            registro = dict(regra)
            if aproximado:
                registro["caminhoNaAbaDeRegras"] = chave
            alvo["regras"].append(registro)

    relatorio = {"regrasSemTag": sorted(sem_tag), "regrasPorAproximacaoDeCaminho": sorted(aproximadas)}
    if not (script and mapeamento and lib):
        return entradas, relatorio

    campos_dataset = ler_mapeamento(mapeamento)
    setters = ler_script(script)
    campos_lib, raizes_sem_confirmacao = ler_lib(lib)

    chaves_minusculas = {chave.lower(): chave for chave in por_chave}
    por_caminho = {}
    for campo, caminho in campos_dataset.items():
        destino = caminho if caminho in por_chave else chaves_minusculas.get(caminho.lower(), caminho)
        por_caminho.setdefault(destino, []).append(campo)

    tx2_ligados = set()
    for entrada in entradas:
        datasets = por_caminho.get(entrada["caminho"] + entrada["tag"], [])
        if not datasets:
            continue
        linhas, tx2, json_campos, json_diretos = [], [], [], []
        for dataset in datasets:
            registro = setters.get(dataset)
            if not registro:
                continue
            linhas.extend(registro["linhas"])
            for campo, dados in registro["fontes"].items():
                tx2_ligados.add(campo)
                if not any(item["campo"] == campo for item in tx2):
                    tx2.append({"campo": campo, "uso": dados["uso"]})
                lib_campo = campos_lib.get(campo, {})
                for caminho in lib_campo.get("json", []):
                    if caminho not in json_campos:
                        json_campos.append(caminho)
                if dados["copiaDireta"]:
                    for caminho in lib_campo.get("copiaDireta", []):
                        if caminho not in json_diretos:
                            json_diretos.append(caminho)
        entrada["plugnotas"] = {
            "dataset": datasets,
            "tx2": tx2,
            "json": json_campos,
            "jsonCopiaDireta": json_diretos,
            "linhasScript": sorted(set(linhas)),
        }

    tx2_lidos = {campo for registro in setters.values() for campo in registro["fontes"]}
    relatorio.update({
        "mapeamentoSemTagNoLeiaute": sorted(
            campo for campo, caminho in campos_dataset.items()
            if caminho not in por_chave and caminho.lower() not in chaves_minusculas),
        "camposDatasetSemScript": sorted(c for c in campos_dataset if c not in setters),
        "tx2DaLibNaoLidosPeloScript": sorted(set(campos_lib) - tx2_lidos),
        "tx2DoScriptSemOrigemNaLib": sorted(tx2_ligados - set(campos_lib)),
        "raizesJsonSemConfirmacao": raizes_sem_confirmacao,
    })
    return entradas, relatorio


# ------------------------------------------------------------------ IBS e CBS

def ler_indop(pasta_anexos):
    livro = openpyxl.load_workbook(pasta_anexos / ANEXO_VII, data_only=True)
    grade = grade_com_mesclagens(livro["cIndOp Public"])
    indicadores = {}
    for linha in grade[1:]:
        codigo = texto(linha[0])
        if not codigo:
            continue
        indicadores[codigo.zfill(6)] = {
            "tipoOperacao": texto(linha[1]),
            "caracteristica": texto(linha[2]),
            "local": texto(linha[3]),
            "dispositivoLegal": texto(linha[4]),
            "observacao": texto(linha[5]),
            "indNFe": texto(linha[6]),
            "indNFSe": texto(linha[7]),
        }
    return indicadores, [texto(valor) for valor in grade[0]]


def ler_correlacao(pasta_anexos):
    livro = openpyxl.load_workbook(pasta_anexos / ANEXO_VIII, data_only=True)
    grade = grade_com_mesclagens(livro["tabela geral"])
    itens, nbs, classificacoes, locais, relacoes, divergencias = {}, {}, {}, [], [], []

    def registrar(dicionario, codigo, descricao, rotulo):
        if not codigo:
            return
        anterior = dicionario.get(codigo)
        if anterior is None:
            dicionario[codigo] = descricao
        elif descricao and anterior != descricao:
            divergencias.append(f"{rotulo} {codigo}: descrições diferentes na planilha")

    for linha in grade[1:]:
        item = texto(linha[0])
        codigo_nbs = texto(linha[2])
        if not item and not codigo_nbs:
            continue
        indop = texto(linha[6])
        indop = indop.split(".")[0].zfill(6) if indop else ""
        classificacao = texto(linha[8])
        classificacao = classificacao.zfill(6) if classificacao.isdigit() else classificacao
        local = texto(linha[7])
        if local and local not in locais:
            locais.append(local)
        registrar(itens, item, texto(linha[1]), "Item")
        registrar(nbs, codigo_nbs, texto(linha[3]), "NBS")
        registrar(classificacoes, classificacao, texto(linha[9]), "cClassTrib")
        relacao = [item, codigo_nbs, texto(linha[4]), texto(linha[5]), indop,
                   locais.index(local) if local else -1, classificacao]
        if relacao not in relacoes:
            relacoes.append(relacao)

    return {"itens": itens, "nbs": nbs, "cClassTrib": classificacoes,
            "locais": locais, "relacoes": relacoes}, divergencias


def ler_regra_inciso_x(pasta_anexos):
    livro = openpyxl.load_workbook(pasta_anexos / ANEXO_VIII, data_only=True)
    return [[texto(valor) for valor in linha] for linha in grade_com_mesclagens(livro["REGRA inc. X"])]


# ------------------------------------------------------------------ execução

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--anexos", required=True, type=Path)
    parser.add_argument("--script", type=Path)
    parser.add_argument("--mapeamento", type=Path)
    parser.add_argument("--lib", type=Path)
    parser.add_argument("--saida", default=Path("definicoes"), type=Path)
    argumentos = parser.parse_args()
    agora = datetime.now(timezone.utc).isoformat(timespec="seconds")

    entradas, relatorio = gerar_de_para(argumentos.anexos, argumentos.script, argumentos.mapeamento, argumentos.lib)
    de_para = {
        "versao": 2,
        "geradoEm": agora,
        "fontes": {
            "leiaute": ANEXO_VI,
            "script": "Scripts/Nacional/LoadEnvio.txt" if argumentos.script else "",
            "mapeamento": "Arquivos/Esquemas/Nacional/v1.01/Mapping.txt" if argumentos.mapeamento else "",
            "lib": "buildTx2/padrao-NACIONAL/props" if argumentos.lib else "",
        },
        "entradas": entradas,
    }
    (argumentos.saida / "de-para-nacional.json").write_text(
        json.dumps(de_para, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

    indicadores, cabecalho = ler_indop(argumentos.anexos)
    correlacao, divergencias = ler_correlacao(argumentos.anexos)
    ibscbs = {
        "versao": 1,
        "geradoEm": agora,
        "fontes": {"indOp": ANEXO_VII, "correlacao": ANEXO_VIII},
        "cabecalhoIndOp": cabecalho,
        "indOp": indicadores,
        **correlacao,
        "regraIncisoX": ler_regra_inciso_x(argumentos.anexos),
    }
    (argumentos.saida / "ibscbs.json").write_text(
        json.dumps(ibscbs, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

    relatorio["ibscbsDivergencias"] = divergencias
    relatorio["indOpDaCorrelacaoAusentesNoAnexoVII"] = sorted(
        {r[4] for r in correlacao["relacoes"] if r[4] and r[4] not in indicadores})

    ligadas = [e for e in entradas if e.get("plugnotas")]
    print(f"de-para: {len(entradas)} tags, {sum(len(e['regras']) for e in entradas)} regras de negócio")
    print(f"         {len(ligadas)} tags ligadas ao script, "
          f"{sum(1 for e in ligadas if e['plugnotas']['json'])} com campo JSON identificado")
    print(f"ibscbs:  {len(indicadores)} indOp, {len(correlacao['itens'])} itens, {len(correlacao['nbs'])} NBS, "
          f"{len(correlacao['cClassTrib'])} cClassTrib, {len(correlacao['relacoes'])} relações")
    Path("ferramentas/relatorio-geracao.json").write_text(
        json.dumps(relatorio, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("relatório salvo em ferramentas/relatorio-geracao.json")


if __name__ == "__main__":
    main()
